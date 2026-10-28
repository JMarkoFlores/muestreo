from datetime import datetime
from io import BytesIO
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import select
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
from app.database import get_db
from app.models import Alert
from app.deps import require_roles
from app.services.audit import audit

router=APIRouter(prefix='/api/reports',tags=['reports'])
def rows(db, severity):
    q=db.query(Alert).order_by(Alert.received_at.desc())
    if severity: q=q.filter(Alert.severity==severity)
    return q.all()
def filename(ext): return f'soc-lite-reporte-{datetime.now().strftime("%Y%m%d-%H%M%S")}.{ext}'
@router.get('/export/{fmt}')
def export(fmt:str,severity:str|None=Query(None),db:Session=Depends(get_db),user=Depends(require_roles('administrador','analista'))):
    if fmt not in {'pdf','docx','xlsx'}: return {'error':'Formato no soportado'}
    data=rows(db,severity); out=BytesIO()
    if fmt=='pdf':
        c=canvas.Canvas(out,pagesize=A4); c.setFont('Helvetica-Bold',14); c.drawString(40,800,'SOC-Lite | Reporte operacional')
        y=770; c.setFont('Helvetica',9)
        for a in data:
            line=f'#{a.id} {a.severity.upper()} {a.source} {a.title[:70]}'
            c.drawString(40,y,line); y-=14
            if y<50: c.showPage(); y=800
        c.save(); media='application/pdf'
    elif fmt=='docx':
        doc=Document(); doc.add_heading('SOC-Lite | Reporte operacional',0); table=doc.add_table(rows=1,cols=5)
        for cell,text in zip(table.rows[0].cells,['ID','Severidad','Origen','Título','Estado']): cell.text=text
        for a in data:
            cells=table.add_row().cells
            for c,text in zip(cells,[str(a.id),a.severity,a.source,a.title,a.status]): c.text=text
        doc.save(out); media='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    else:
        wb=Workbook(); ws=wb.active; ws.title='Alertas'; ws.append(['ID','Severidad','Origen','Título','Estado','Recibida'])
        for a in data: ws.append([a.id,a.severity,a.source,a.title,a.status,a.received_at.isoformat()])
        wb.save(out); media='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    out.seek(0); audit(db,user.username,'reports','export','ok','report',fmt,details={'severity':severity,'rows':len(data)})
    return StreamingResponse(out,media_type=media,headers={'Content-Disposition':f'attachment; filename={filename(fmt)}'})
